import importlib.util
from pathlib import Path
import tempfile
import unittest
import zipfile
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location('memory_extract', ROOT/'skills/rc-documents/scripts/extract_memory.py')
extract = importlib.util.module_from_spec(spec)
spec.loader.exec_module(extract)


def pdf(path, texts):
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
               ('<< /Type /Pages /Kids ['+' '.join(str(4+2*i)+' 0 R' for i in range(len(texts)))+'] /Count '+str(len(texts))+' >>').encode(),
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    for i,text in enumerate(texts):
        objects.append(('<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents '+str(5+2*i)+' 0 R >>').encode())
        content=('BT /F1 12 Tf 72 720 Td ('+text+') Tj ET').encode()
        objects.append(('<< /Length '+str(len(content))+' >>\nstream\n').encode()+content+b'\nendstream')
    data=b'%PDF-1.4\n'; offsets=[0]
    for number,obj in enumerate(objects,1):
        offsets.append(len(data));data+=str(number).encode()+b' 0 obj\n'+obj+b'\nendobj\n'
    xref=len(data)
    data+=('xref\n0 '+str(len(offsets))+'\n0000000000 65535 f \n').encode()
    data+=b''.join(('%010d 00000 n \n'%offset).encode() for offset in offsets[1:])
    data+=('trailer\n<< /Size '+str(len(offsets))+' /Root 1 0 R >>\nstartxref\n'+str(xref)+'\n%%EOF').encode()
    path.write_bytes(data)


class MemoryExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):
        self.temp.cleanup()
    def test_docx_headings_paragraphs_and_table_cells(self):
        path=self.root/'proposal.docx';document=Document();document.add_heading('Commercial terms',1);document.add_paragraph('Keep the proposal editable.');table=document.add_table(rows=2,cols=2);table.cell(0,0).text='Deliverable';table.cell(0,1).text='Cost';table.cell(1,0).text='Unicode proposal — café';table.cell(1,1).text='IDR 25,000,000';document.save(path)
        original=path.read_bytes();result=extract.Extractor({}).run(path,'docx');self.assertEqual(result['state'],'ready');self.assertEqual(path.read_bytes(),original)
        cell=next(p for p in result['passages'] if '25,000,000' in p['text']);self.assertEqual(cell['locator']['kind'],'table-cell');self.assertEqual((cell['locator']['table'],cell['locator']['row'],cell['locator']['cell']),(1,2,2));self.assertEqual(cell['locator']['heading'],'Commercial terms');self.assertNotIn('page',cell['locator'])
    def test_pdf_page_citations_and_page_limit(self):
        path=self.root/'proposal.pdf';pdf(path,['FirstPageEvidence','SecondPageEvidence']);result=extract.Extractor({}).run(path,'pdf');self.assertEqual(result['state'],'ready');self.assertEqual([p['locator']['page'] for p in result['passages']],[1,2]);self.assertIn('SecondPageEvidence',result['passages'][1]['text'])
        limited=extract.Extractor({'maxPages':1}).run(path,'pdf');self.assertEqual(limited['state'],'partial');self.assertEqual(len(limited['passages']),1)
    def test_empty_pdf_requires_ocr(self):
        path=self.root/'scan.pdf';pdf(path,['']);result=extract.Extractor({}).run(path,'pdf');self.assertEqual(result['state'],'needs_ocr');self.assertEqual(result['passages'],[])
    def test_text_boundaries_and_explicit_partial_coverage(self):
        path=self.root/'source.txt';text='a'*65+'BoundaryToken'+'z'*100;path.write_text(text);result=extract.Extractor({'passageChars':80,'overlapChars':30}).run(path,'txt');self.assertTrue(any('BoundaryToken' in p['text'] for p in result['passages']));self.assertTrue(all(p['locator']['startLine']==1 for p in result['passages']))
        limited=extract.Extractor({'maxTextBytes':40}).run(path,'txt');self.assertEqual(limited['state'],'partial');self.assertLessEqual(limited['textBytes'],40)
    def test_office_expansion_and_entities_fail_without_external_reads(self):
        path=self.root/'bad.docx'
        with zipfile.ZipFile(path,'w') as archive:archive.writestr('word/document.xml','<!DOCTYPE test [<!ENTITY secret SYSTEM "file:///etc/passwd">]><test>&secret;</test>')
        with self.assertRaisesRegex(ValueError,'entities'):extract.Extractor({}).run(path,'docx')
        with self.assertRaisesRegex(ValueError,'expansion'):extract.Extractor({'maxExpandedBytes':10}).run(path,'docx')
    def test_empty_and_unknown_formats_never_report_ready(self):
        path=self.root/'empty.txt';path.write_text('  ');self.assertEqual(extract.Extractor({}).run(path,'txt')['state'],'failed');self.assertEqual(extract.Extractor({}).run(path,'xlsx')['state'],'unsupported')


if __name__=='__main__':unittest.main()
