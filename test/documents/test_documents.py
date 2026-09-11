import base64
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest
import zipfile
from docx import Document
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("documents", ROOT / "skills/rc-documents/scripts/documents.py")
documents = importlib.util.module_from_spec(spec)
spec.loader.exec_module(documents)


class DocumentTests(unittest.TestCase):
    def setUp(self):
        output = Path(os.environ.get("X056_DOCUMENT_TEST_OUTPUT", ROOT / ".deploy/chat-document-tests"))
        output.mkdir(parents=True, exist_ok=True)
        self.directory = Path(tempfile.mkdtemp(dir=output))
        self.original = self.directory / "proposal.docx"
        document = Document()
        document.add_heading("Proposal", 0)
        paragraph = document.add_paragraph()
        paragraph.add_run("Original ").bold = True
        paragraph.add_run("proposal")
        paragraph.add_run(" for review.").italic = True
        document.add_paragraph("Retain numbered scope", style="List Number")
        table = document.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Cost"
        table.cell(1, 0).text = "Service"
        table.cell(1, 1).text = "100"
        document.sections[0].header.paragraphs[0].text = "Proposal header"
        document.sections[0].footer.paragraphs[0].text = "Proposal footer"
        document.add_page_break()
        document.add_paragraph("Delivery schedule")
        image = self.directory / "pixel.png"
        image.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="))
        document.add_picture(str(image))
        document.save(self.original)

    def test_preserves_unaffected_structures_and_rendered_pages(self):
        before = documents.inspect(self.original)
        baseline = documents.preview(self.original, self.directory / "before")
        revised = self.directory / "revised.docx"
        revised.write_bytes(self.original.read_bytes())
        result = documents.edit(revised, {"replacements": [{"find": "Original proposal", "replace": "Revised proposal"}]})
        after = documents.inspect(revised)
        self.assertEqual(result["changedParts"], ["word/document.xml"])
        self.assertEqual(before["structures"], after["structures"])
        for name, checksum in before["parts"].items():
            if name != "word/document.xml":
                self.assertEqual(after["parts"][name], checksum, name)
        doc = Document(revised)
        self.assertTrue(doc.paragraphs[1].runs[0].bold)
        self.assertTrue(doc.paragraphs[1].runs[2].italic)
        self.assertEqual(doc.tables[0].cell(1, 1).text, "100")
        self.assertEqual(doc.paragraphs[1].text, "Revised proposal for review.")
        rendered = documents.preview(revised, self.directory / "after")
        self.assertEqual(len(baseline["pages"]), 2)
        self.assertEqual(len(rendered["pages"]), 2)
        self.assertEqual(Path(baseline["pages"][1]).read_bytes(), Path(rendered["pages"][1]).read_bytes())
        self.assertNotEqual(Path(baseline["pages"][0]).read_bytes(), Path(rendered["pages"][0]).read_bytes())
        self.assertIn("Original proposal", json.dumps(documents.inspect(self.original)))

    def test_refuses_ambiguous_or_unsupported_edits_without_touching_file(self):
        original = self.original.read_bytes()
        with self.assertRaisesRegex(ValueError, "Expected"):
            documents.edit(self.original, {"replacements": [{"find": "absent", "replace": "wrong"}]})
        self.assertEqual(self.original.read_bytes(), original)
        doc = Document(self.original)
        field = OxmlElement("w:fldChar")
        doc.paragraphs[1]._p.append(field)
        doc.save(self.original)
        original = self.original.read_bytes()
        with self.assertRaisesRegex(ValueError, "unsupported structure"):
            documents.edit(self.original, {"replacements": [{"find": "Original proposal", "replace": "Wrong"}]})
        self.assertEqual(self.original.read_bytes(), original)

    def test_creation_and_archive_limits(self):
        created = self.directory / "created.docx"
        documents.create(created, {"title": "Proposal", "paragraphs": ["Scope"], "tables": [[["Cost"], ["100"]]]})
        self.assertEqual(Document(created).tables[0].cell(1, 0).text, "100")
        bad = self.directory / "bad.zip"
        with zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("huge.txt", b"0" * 1024 * 1024)
        with self.assertRaisesRegex(ValueError, "expansion"):
            documents.package(bad)
        with zipfile.ZipFile(bad, "w") as archive:
            archive.writestr("../escape", "bad")
        with self.assertRaisesRegex(ValueError, "Unsafe"):
            documents.package(bad)

    def test_preview_refuses_external_resources_without_changing_the_document(self):
        with zipfile.ZipFile(self.original, "a") as archive:
            archive.writestr("word/_rels/external.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="external" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" TargetMode="External" Target="http://127.0.0.1/private"/></Relationships>')
        original = self.original.read_bytes()
        with self.assertRaisesRegex(ValueError, "externally linked"):
            documents.preview(self.original, self.directory / "blocked")
        self.assertEqual(self.original.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
